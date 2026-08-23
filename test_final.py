import requests
from docx import Document
import os

API_URL = "http://localhost:8000"

def create_sample_docx(filename):
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("This is a test paragraph.")
    doc.save(filename)

def verify_final_features():
    filename = "final_test.docx"
    create_sample_docx(filename)
    
    print("Testing Final Features...")
    
    # 1. Test LaTeX Export
    with open(filename, "rb") as f:
        files = {"file": f}
        data = {"style": "IEEE", "output_format": "latex"}
        response = requests.post(f"{API_URL}/process", files=files, data=data)
        
    assert response.status_code == 200, f"LaTeX Request failed: {response.text}"
    json_data = response.json()
    download_url = json_data["download_url"]
    server_filename = download_url.split("/")[-1] # Get the formatted filename
    
    assert download_url.endswith(".tex"), f"Expected .tex output, got {download_url}"
    print("LaTeX Export: OK")
    
    # 2. Test Preview Endpoint
    # We need to use the actual filename on server (formatted_...)
    preview_response = requests.get(f"{API_URL}/preview/{server_filename}")
    assert preview_response.status_code == 200, f"Preview failed: {preview_response.text}"
    content = preview_response.json()["content"]
    assert "This is a test paragraph" in content, "Preview content mismatch"
    print("Preview Endpoint: OK")

    print("\nAll final features verified!")
    
    if os.path.exists(filename):
        os.remove(filename)

if __name__ == "__main__":
    verify_final_features()
