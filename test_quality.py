import requests
from docx import Document
import os

API_URL = "http://localhost:8000"

def create_bad_order_docx(filename):
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Intro text.")
    doc.add_heading("Abstract", level=1) # Wrong order
    doc.add_paragraph("Abstract text.")
    doc.save(filename)

def verify_quality_checks():
    filename = "quality_test.docx"
    create_bad_order_docx(filename)
    
    print("Testing Quality Checks...")
    with open(filename, "rb") as f:
        files = {"file": f}
        data = {"style": "Standard"}
        response = requests.post(f"{API_URL}/process", files=files, data=data)
        
    assert response.status_code == 200, f"Request failed: {response.text}"
    json_data = response.json()
    
    assert "quality_report" in json_data, "Missing quality_report"
    report = json_data["quality_report"]
    
    # Check Section Order
    order_check = next((item for item in report if item["rule"] == "Section Order"), None)
    assert order_check is not None, "Section Order rule missing"
    assert order_check["status"] == "fail", f"Section Order should fail: {order_check}"
    print("Section Order Check: OK (Correctly failed)")
    
    # Check Missing Sections (Conclusion, References missing)
    missing_check = next((item for item in report if item["rule"] == "Missing Sections"), None)
    assert missing_check is not None, "Missing Sections rule missing"
    assert "Conclusion" in missing_check["message"], "Conclusion should be missing"
    print("Missing Sections Check: OK")
    
    # Check Keywords (Abstract exists, so should try to extract)
    keyword_check = next((item for item in report if item["rule"] == "Smart Keywords"), None)
    assert keyword_check is not None, "Smart Keywords rule missing"
    # Note: extraction might fail if text is too short, but rule should exist
    print(f"Smart Keywords Check: OK ({keyword_check['message']})")

    print("\nAll quality checks passed!")
    
    if os.path.exists(filename):
        os.remove(filename)

if __name__ == "__main__":
    verify_quality_checks()
