
import requests
import os
import time

BASE_URL = "http://localhost:8003"

def test_chat_flow():
    print("--- Testing Phase 7: Chat Assistant (TF-IDF mode) ---")
    
    # 1. Create a dummy test doc
    filename = "test_chat_doc_v2.docx"
    if not os.path.exists(filename):
        from docx import Document
        doc = Document()
        doc.add_heading('Introduction to ManuScript AI', 0)
        doc.add_paragraph("ManuScript is an AI-powered academic assistant.")
        doc.add_paragraph("It used to use embeddings but now uses TF-IDF for speed.")
        doc.save(filename)
        print(f"Created temporary file: {filename}")

    # 2. Upload
    print("Uploading file to index (should be fast)...")
    with open(filename, 'rb') as f:
        files = {'file': f}
        data = {
            'style': 'Standard',
            'output_format': 'docx',
            'add_signature': 'false',
            'use_sample': 'false'
        }
        try:
            resp = requests.post(f"{BASE_URL}/process", files=files, data=data, timeout=30)
        except requests.Timeout:
             print("❌ Upload Timed Out!")
             return

    if resp.status_code == 200:
        result = resp.json()
        download_url = result.get("download_url")
        target_filename = download_url.split("/")[-1]
        print(f"Upload successful. Target filename: {target_filename}")
        
        # 3. Chat Query
        query = "What does it use for speed?"
        print(f"Querying: '{query}'")
        
        chat_resp = requests.post(f"{BASE_URL}/chat", json={
            "filename": target_filename,
            "query": query
        })
        
        if chat_resp.status_code == 200:
            answer = chat_resp.json().get("answer")
            print(f"✅ Chat Answer: {answer}")
            if "TF-IDF" in answer:
                print("   (Correctly retrieved context!)")
        else:
            print(f"❌ Chat Failed: {chat_resp.status_code} - {chat_resp.text}")

    else:
        print(f"❌ Upload Failed: {resp.status_code} - {resp.text}")

if __name__ == "__main__":
    test_chat_flow()
