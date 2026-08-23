
import requests
import os
import time

BASE_URL = "http://localhost:8000"

def test_chat_flow():
    print("--- Testing Phase 7: Chat Assistant ---")
    
    # 1. Create a dummy test doc
    filename = "test_chat_doc.docx"
    if not os.path.exists(filename):
        from docx import Document
        doc = Document()
        doc.add_heading('Introduction to ManuScript AI', 0)
        doc.add_paragraph("ManuScript is an AI-powered academic assistant.")
        doc.add_paragraph("It features a new Chat Assistant that uses RAG technology.")
        doc.add_paragraph("The system uses hybrid retrieval with TF-IDF and embeddings.")
        doc.save(filename)
        print(f"Created temporary file: {filename}")

    # 2. Upload
    print("Uploading file to index...")
    with open(filename, 'rb') as f:
        files = {'file': f}
        data = {
            'style': 'Standard',
            'output_format': 'docx',
            'add_signature': 'false',
            'use_sample': 'false'
        }
        resp = requests.post(f"{BASE_URL}/process", files=files, data=data)
        
    if resp.status_code == 200:
        result = resp.json()
        download_url = result.get("download_url")
        # Extract filename expected by chat (formatted one usually)
        target_filename = download_url.split("/")[-1]
        print(f"Upload successful. Target filename: {target_filename}")
        
        # 3. Chat Query
        query = "What technology does the Chat Assistant use?"
        print(f"Querying: '{query}'")
        
        chat_resp = requests.post(f"{BASE_URL}/chat", json={
            "filename": target_filename,
            "query": query
        })
        
        if chat_resp.status_code == 200:
            answer = chat_resp.json().get("answer")
            print(f"✅ Chat Answer: {answer}")
            if "RAG" in answer:
                print("   (Correctly retrieved context!)")
            else:
                print("   (Warning: Context might be missing or generic fallback)")
        else:
            print(f"❌ Chat Failed: {chat_resp.status_code} - {chat_resp.text}")

    else:
        print(f"❌ Upload Failed: {resp.status_code} - {resp.text}")

if __name__ == "__main__":
    try:
        test_chat_flow()
    except Exception as e:
        print(f"Test crashed: {e}")
