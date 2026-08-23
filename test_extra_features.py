import requests
from docx import Document
import os
import time

API_URL = "http://localhost:8003"

def create_test_docx(filename):
    doc = Document()
    doc.add_paragraph("Test document for extra features.")
    doc.save(filename)

def verify_extra_features():
    filename = "extra_test.docx"
    create_test_docx(filename)
    
    # Test 1: Signature Line
    print("Testing Signature Line...")
    with open(filename, "rb") as f:
        files = {"file": f}
        data = {
            "style": "Standard",
            "add_signature": "true",
            "output_format": "docx"
        }
        response = requests.post(f"{API_URL}/process", files=files, data=data)
        
    assert response.status_code == 200, f"Request failed: {response.text}"
    json_data = response.json()
    download_url = json_data.get("download_url")
    
    print("Downloading Signature output...")
    output_response = requests.get(download_url)
    with open("output_signature.docx", "wb") as f:
        f.write(output_response.content)
        
    doc_sig = Document("output_signature.docx")
    # Check last paragraphs for signature
    last_para = doc_sig.paragraphs[-1].text.strip()
    second_last = doc_sig.paragraphs[-2].text.strip()
    
    print(f"Last paragraph: '{last_para}'")
    print(f"Second last: '{second_last}'")
    
    assert "Signature" in last_para, "Signature text not found"
    assert "_" in second_last, "Signature line not found"
    print("Signature Line: OK")
    
    # Test 2: PDF Conversion
    print("\nTesting PDF Conversion...")
    with open(filename, "rb") as f:
        files = {"file": f}
        data = {
            "style": "Standard",
            "add_signature": "false",
            "output_format": "pdf"
        }
        response = requests.post(f"{API_URL}/process", files=files, data=data)
        
    json_data = response.json()
    download_url = json_data.get("download_url")
    stats = json_data.get("stats", {})
    
    if "warning" in stats:
        print(f"PDF Conversion Warning: {stats['warning']}")
        print("Skipping PDF verification (likely no Word installed).")
    else:
        print(f"Download URL: {download_url}")
        assert download_url.endswith(".pdf"), "Download URL does not end with .pdf"
        
        # Try downloading
        print("Downloading PDF output...")
        output_response = requests.get(download_url)
        assert output_response.status_code == 200, "Failed to download PDF"
        assert output_response.headers["content-type"] == "application/pdf", "Content-Type is not PDF"
        
        with open("output.pdf", "wb") as f:
            f.write(output_response.content)
        print("PDF Conversion: OK")

    print("\nAll extra feature tests passed!")
    
    # Cleanup
    if os.path.exists(filename):
        os.remove(filename)
    if os.path.exists("output_signature.docx"):
        os.remove("output_signature.docx")
    if os.path.exists("output.pdf"):
        os.remove("output.pdf")

if __name__ == "__main__":
    verify_extra_features()
