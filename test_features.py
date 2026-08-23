import requests
from docx import Document
from docx.shared import Pt
import os

API_URL = "http://localhost:8000"

def create_test_docx(filename):
    doc = Document()
    doc.add_paragraph("This is a test document.")
    doc.add_paragraph("It has a few sentences to calculate readability.")
    doc.add_paragraph("We expect the word count to be around 20.")
    doc.save(filename)

def verify_features():
    filename = "feature_test.docx"
    create_test_docx(filename)
    
    # Test 1: Standard Style & Stats
    print("Testing Standard Style & Stats...")
    with open(filename, "rb") as f:
        files = {"file": f}
        data = {"style": "Standard"}
        response = requests.post(f"{API_URL}/process", files=files, data=data)
        
    assert response.status_code == 200, f"Request failed: {response.text}"
    json_data = response.json()
    
    # Check Stats
    stats = json_data.get("stats")
    assert stats is not None, "Stats missing in response"
    print(f"Stats received: {stats}")
    assert stats["word_count"] > 10, "Word count too low"
    assert "readability_score" in stats, "Readability score missing"
    
    # Check Download URL
    download_url = json_data.get("download_url")
    assert download_url is not None, "Download URL missing"
    
    # Download and check formatting (Standard = 1.5 spacing)
    print("Downloading Standard output...")
    output_response = requests.get(download_url)
    with open("output_standard.docx", "wb") as f:
        f.write(output_response.content)
        
    # Check Body Text (Index 2 likely, as 0 is Title, 1 might be Author/Body)
    # In create_test_docx:
    # 0: "This is a test document." -> Title
    # 1: "It has a few sentences..." -> Body
    # 2: "We expect..." -> Body
    
    doc_std = Document("output_standard.docx")
    spacing = doc_std.paragraphs[1].paragraph_format.line_spacing
    print(f"Standard Line Spacing (Body): {spacing}")
    assert spacing == 1.5, f"Expected 1.5 spacing, got {spacing}"
    
    # Test 2: APA Style (Double Spacing)
    print("\nTesting APA Style...")
    with open(filename, "rb") as f:
        files = {"file": f}
        data = {"style": "APA"}
        response = requests.post(f"{API_URL}/process", files=files, data=data)
        
    json_data = response.json()
    download_url = json_data.get("download_url")
    
    print("Downloading APA output...")
    output_response = requests.get(download_url)
    with open("output_apa.docx", "wb") as f:
        f.write(output_response.content)
        
    doc_apa = Document("output_apa.docx")
    spacing = doc_apa.paragraphs[1].paragraph_format.line_spacing
    print(f"APA Line Spacing (Body): {spacing}")
    assert spacing == 2.0, f"Expected 2.0 spacing, got {spacing}"
    
    print("\nAll feature tests passed!")
    
    # Cleanup
    if os.path.exists(filename):
        os.remove(filename)
    if os.path.exists("output_standard.docx"):
        os.remove("output_standard.docx")
    if os.path.exists("output_apa.docx"):
        os.remove("output_apa.docx")

if __name__ == "__main__":
    verify_features()
