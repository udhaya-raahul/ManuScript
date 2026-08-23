
import os
import requests
import json
from docx import Document
from docx.shared import Inches

BASE_URL = "http://localhost:8000"

def create_test_docx(filename):
    doc = Document()
    doc.add_heading("Deep Learning for Medical Imaging Analysis", 0)
    
    doc.add_heading("Abstract", 1)
    doc.add_paragraph("This paper explores the use of machine learning and neural networks for analyzing medical images. We propose a new algorithm for tumor detection.")
    
    doc.add_heading("Introduction", 1)
    doc.add_paragraph("Artificial Intelligence (AI) has revolutionized healthcare. As seen in Figure 1, the growth is exponential.")
    
    # Add a dummy image (we can't easily add a real image without a file, skipping image add for now, 
    # just checking text triggers. Wait, if no image added, auditor returns nothing or warn?
    # Logic: if image_count > 0 and no captions.
    # If I don't add image, I can't test "missing caption" warning easily without an image file.
    # But I can test Presentation and Audio generation without images.
    
    doc.save(filename)
    return filename

def test_phase6_features():
    print("Testing Phase 6: Beyond the Paper...")
    
    filename = "test_phase6.docx"
    create_test_docx(filename)
    
    try:
        with open(filename, "rb") as f:
            files = {"file": (filename, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            data = {"style": "IEEE", "add_signature": "false"}
            
            print(f"Uploading {filename}...")
            response = requests.post(f"{BASE_URL}/process", files=files, data=data)
            
            if response.status_code == 200:
                result = response.json()
                download_url = result["download_url"]
                output_filename = download_url.split("/")[-1]
                print(f"✅ Upload Success. Output: {output_filename}")
                
                # 1. Test Slides Generation
                print("Testing Slides Generation...")
                slides_resp = requests.post(f"{BASE_URL}/generate_slides/{output_filename}")
                if slides_resp.status_code == 200:
                    print(f"✅ Slides Generated: {slides_resp.json()['download_url']}")
                else:
                    print(f"❌ Slides Failed: {slides_resp.text}")

                # 2. Test Audio Generation
                print("Testing Audio Generation...")
                audio_resp = requests.post(f"{BASE_URL}/generate_audio/{output_filename}")
                if audio_resp.status_code == 200:
                    print(f"✅ Audio Generated: {audio_resp.json()['download_url']}")
                else:
                    print(f"❌ Audio Failed: {audio_resp.text}")
                    
                # 3. Check Figure Auditor (Expect Pass as 0 images)
                report = result["quality_report"]
                fig_audit = [item for item in report if item["rule"] == "Figure Auditor"]
                if fig_audit:
                    print(f"✅ Figure Auditor: {fig_audit[0]['status']}")
                else:
                    print("❌ Figure Auditor rule missing")

            else:
                print(f"❌ Upload Failed: {response.status_code} - {response.text}")
                
    except Exception as e:
        print(f"❌ Error: {e}")
    finally:
        if os.path.exists(filename):
            os.remove(filename)

if __name__ == "__main__":
    test_phase6_features()
