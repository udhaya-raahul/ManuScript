from docx import Document
import os

def create_informal_docx(filename):
    doc = Document()
    doc.add_heading("My Awesome Research", level=1)
    
    doc.add_heading("Abstract", level=2)
    doc.add_paragraph("This paper is basically about some stuff we did with kids in a lab. It was very good because we got a lot of data. The results were awesome and not bad at all.")
    
    doc.add_heading("Introduction", level=2)
    doc.add_paragraph("A lot of researchers think that kids are hard to study. Basically, we found that stuff is easy if you have a good guy helping you.")
    
    doc.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    create_informal_docx("informal_test.docx")
