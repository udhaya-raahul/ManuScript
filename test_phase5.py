
import os
import requests
import json
from docx import Document

BASE_URL = "http://localhost:8000"

def create_test_docx(filename):
    doc = Document()
    doc.add_heading("Deep Learning for Medical Imaging Analysis", 0)
    
    doc.add_heading("Abstract", 1)
    doc.add_paragraph("This paper explores the use of machine learning and neural networks for analyzing medical images. We propose a new algorithm for tumor detection.")
    
    doc.add_heading("Introduction", 1)
    doc.add_paragraph("Artificial Intelligence (AI) has revolutionized healthcare. We focus on deep learning techniques.")
    
    doc.add_heading("References", 1)
    doc.add_paragraph("Vaswani, A. et al. Attention Is All You Need. NIPS 2017.")
    doc.add_paragraph("Devlin, J. BERT: Pre-training of Deep Bidirectional Transformers.")
    
    doc.save(filename)
    return filename

def test_phase5_features():
    print("Testing Phase 5: Advanced Researcher Suite...")
    
    filename = "test_phase5.docx"
    create_test_docx(filename)
    
    try:
        with open(filename, "rb") as f:
            files = {"file": (filename, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            data = {"style": "IEEE", "add_signature": "false"}
            
            print(f"Uploading {filename}...")
            response = requests.post(f"{BASE_URL}/process", files=files, data=data)
            
            if response.status_code == 200:
                result = response.json()
                
                if "quality_report" not in result:
                    print("❌ 'quality_report' MISSING in response!")
                    print(f"Response keys: {result.keys()}")
                    return
                
                report = result["quality_report"]
                
                # Check Journal Recommendations
                journals = [item for item in report if item["rule"] == "Journal Recommendations"]
                if journals:
                    print("✅ Journal Recommendations: PASS")
                    print(f"   Top Match: {journals[0]['details'][0]['name']}")
                else:
                    print("❌ Journal Recommendations: FAIL (Not found)")
                    
                # Check Bibliography Linker
                bibs = [item for item in report if item["rule"] == "Bibliography Linker"]
                if bibs:
                    print("✅ Bibliography Linker: PASS")
                    print(f"   Found {bibs[0]['message']}")
                    print(f"   Sample Link: {bibs[0]['details'][0]['url']}")
                else:
                    print("❌ Bibliography Linker: FAIL (Not found)")
                    
            else:
                print(f"❌ Upload Failed: {response.status_code} - {response.text}")
                
    except Exception as e:
        print(f"❌ Error: {e}")
    finally:
        if os.path.exists(filename):
            os.remove(filename)

if __name__ == "__main__":
    test_phase5_features()
