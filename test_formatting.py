from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend.formatter import format_manuscript
import os

def create_messy_docx(filename):
    doc = Document()
    
    # Title (messy)
    p = doc.add_paragraph("  My Great Manuscript  ")
    p.runs[0].font.size = Pt(14)
    p.runs[0].bold = True
    
    # Author
    doc.add_paragraph("By Author Name")
    
    # Heading 1 (messy)
    p = doc.add_paragraph("Introduction")
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(13)
    
    # Body text (messy)
    p = doc.add_paragraph("This is the introduction. It has some random formatting.")
    p.runs[0].font.name = "Arial"
    
    # Heading 2 (messy)
    p = doc.add_paragraph("Methods")
    p.runs[0].bold = True
    
    # Body text
    doc.add_paragraph("We did some things. It was great.")
    
    doc.save(filename)

def verify_formatting(filename):
    doc = Document(filename)
    
    print(f"Verifying {filename}...")
    
    # Check Margins
    section = doc.sections[0]
    assert section.top_margin == Inches(1), "Top margin incorrect"
    assert section.bottom_margin == Inches(1), "Bottom margin incorrect"
    print("Margins: OK")
    
    # Check Title (First paragraph)
    title = doc.paragraphs[0]
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER, "Title alignment incorrect"
    # Note: We re-added runs, so check the first run
    assert title.runs[0].font.name == 'Times New Roman', "Title font incorrect"
    assert title.runs[0].font.size == Pt(16), "Title size incorrect"
    print("Title: OK")
    
    # Check Heading (Introduction)
    # Index 0=Title, 1=Author (Body), 2=Introduction (Heading)
    # Wait, Author might be body text now based on logic? 
    # Logic: Title is i=0. i=1 is Author. 
    # Author "By Author Name" -> Not bold, not >12pt -> Body Text.
    # Introduction "Introduction" -> Bold, 13pt -> Heading.
    
    intro = doc.paragraphs[2]
    assert intro.text.strip() == "Introduction"
    assert intro.alignment == WD_ALIGN_PARAGRAPH.LEFT, "Heading alignment incorrect"
    assert intro.runs[0].font.size == Pt(14), "Heading size incorrect"
    print("Heading 1: OK")
    
    # Check Body Text
    body = doc.paragraphs[3]
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY, "Body alignment incorrect"
    assert body.runs[0].font.name == 'Times New Roman', "Body font incorrect"
    assert body.runs[0].font.size == Pt(12), "Body size incorrect"
    print("Body Text: OK")
    
    print("All checks passed!")

if __name__ == "__main__":
    input_file = "messy_input.docx"
    output_file = "clean_output.docx"
    
    create_messy_docx(input_file)
    format_manuscript(input_file, output_file)
    verify_formatting(output_file)
    
    # Cleanup
    if os.path.exists(input_file):
        os.remove(input_file)
    if os.path.exists(output_file):
        os.remove(output_file)
